"""Excel export (openpyxl) for every main table + the import template."""

import io
import re

from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from lib.db import db
from lib.dates import today_iso
from lib.helpers import build_kartu, fdate, log_action
from lib.security import get_current_user, require_admin
from routers.items import item_query

router = APIRouter(prefix="/export", tags=["export"])

THIN = Side(style="thin", color="D1D5DB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


async def _kop() -> tuple[str, str]:
    p = await db.school_profile.find_one({"id": "profil"}) or {}
    kop1 = p.get("nama_sekolah") or "SIMBARA SEKOLAH"
    parts = [
        f"NPSN: {p['npsn']}" if p.get("npsn") else "",
        p.get("alamat", ""),
        ", ".join(x for x in [p.get("kota", ""), p.get("provinsi", "")] if x),
        f"Telp: {p['telepon']}" if p.get("telepon") else "",
    ]
    return kop1, " | ".join(x for x in parts if x)


def _sheet(
    wb: Workbook,
    name: str,
    kop1: str,
    kop2: str,
    headers: list[str],
    rows: list[list[object]],
    widths: list[int],
) -> None:
    if wb.sheetnames == ["Sheet"]:
        ws = wb.active
        ws.title = name
    else:
        ws = wb.create_sheet(name)
    ncols = len(headers)
    last = get_column_letter(ncols)
    ws.merge_cells(f"A1:{last}1")
    c1 = ws.cell(row=1, column=1, value=kop1)
    c1.font = Font(bold=True, size=13, color="064E3B")
    c1.alignment = Alignment(horizontal="center")
    ws.row_dimensions[1].height = 22
    ws.merge_cells(f"A2:{last}2")
    c2 = ws.cell(row=2, column=1, value=kop2)
    c2.font = Font(size=9, color="475569")
    c2.alignment = Alignment(horizontal="center")
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col, value=h)
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor="065F46")
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[4].height = 20
    for r, row in enumerate(rows, start=5):
        for col, val in enumerate(row, start=1):
            cell = ws.cell(row=r, column=col, value=val)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="center")
            if isinstance(val, (int, float)) and not isinstance(val, bool):
                cell.number_format = "#,##0"
    for col, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = "A5"


def _response(wb: Workbook, filename: str) -> StreamingResponse:
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/{jenis}")
async def export_excel(jenis: str, request: Request, user: dict = Depends(get_current_user)):
    def qp(key: str) -> str:
        return (request.query_params.get(key) or "").strip()

    kop1, kop2 = await _kop()
    wb = Workbook()
    stamp = today_iso().replace("-", "")
    label = ""

    if jenis in ("barang", "persediaan", "inventaris"):
        preset = {"persediaan": "persediaan", "inventaris": "inventaris"}.get(jenis, "")
        docs = await db.items.find(
            item_query(qp("q"), preset or qp("jenis"), qp("kategori"), qp("kondisi"), qp("lokasi"), qp("status"), qp("tahun"))
        ).sort([("kode", 1)]).to_list(50000)
        rows = [
            [
                n + 1,
                d.get("kode", ""),
                d.get("nama", ""),
                d.get("kategori", ""),
                "Inventaris/Aset" if d.get("jenis") == "inventaris" else "Persediaan",
                d.get("merk", ""),
                d.get("tipe_model", ""),
                d.get("spesifikasi", ""),
                d.get("satuan", ""),
                d.get("tahun") or "",
                d.get("sumber_dana", ""),
                d.get("harga_satuan", 0),
                d.get("stok", 0),
                d.get("lokasi", ""),
                d.get("kondisi", ""),
                "Arsip" if d.get("deleted_at") else "Aktif",
                d.get("nomor_seri", ""),
                d.get("nup", ""),
                d.get("keterangan", ""),
            ]
            for n, d in enumerate(docs)
        ]
        headers = ["No", "Kode Barang", "Nama Barang", "Kategori", "Jenis", "Merk", "Tipe", "Spesifikasi", "Satuan",
                   "Tahun Perolehan", "Sumber Dana", "Harga (Rp)", "Stok", "Lokasi", "Kondisi", "Status",
                   "Nomor Seri", "NUP", "Keterangan"]
        _sheet(wb, "Data Barang", kop1, kop2, headers, rows,
               [5, 12, 30, 16, 15, 14, 14, 24, 9, 12, 14, 14, 9, 18, 13, 10, 16, 10, 26])
        fname = f"Data-Barang-{stamp}.xlsx"
        label = "Data Barang"

    elif jenis in ("masuk", "keluar", "penyesuaian"):
        query: dict = {"jenis": jenis, "status": "active"}
        if qp("tanggal_awal"):
            query["tanggal"] = {**query.get("tanggal", {}), "$gte": qp("tanggal_awal")}
        if qp("tanggal_akhir"):
            query["tanggal"] = {**query.get("tanggal", {}), "$lte": qp("tanggal_akhir")}
        if qp("kategori"):
            query["kategori"] = qp("kategori")
        if qp("q"):
            rx = {"$regex": re.escape(qp("q")), "$options": "i"}
            query["$or"] = [{"nomor": rx}, {"nama_barang": rx}, {"penerima": rx}, {"keterangan": rx}]
        docs = await db.item_transactions.find(query).sort([("tanggal", -1), ("created_at", -1)]).to_list(50000)
        if jenis == "masuk":
            rows = [
                [
                    i + 1,
                    t.get("nomor", ""),
                    fdate(t.get("tanggal", "")),
                    t.get("kode_barang", ""),
                    t.get("nama_barang", ""),
                    t.get("kategori", ""),
                    "Barang Masuk",
                    t.get("jumlah", 0),
                    t.get("satuan", ""),
                    t.get("harga_satuan", 0),
                    t.get("total_harga", 0),
                    t.get("lokasi", ""),
                    t.get("sumber_dana", ""),
                    t.get("nomor_dokumen", ""),
                    t.get("kondisi", ""),
                    t.get("keterangan", ""),
                    t.get("user_name", ""),
                ]
                for i, t in enumerate(docs)
            ]
            headers = ["No", "Nomor Transaksi", "Tanggal", "Kode Barang", "Nama Barang", "Kategori", "Jenis Transaksi",
                       "Jumlah", "Satuan", "Harga (Rp)", "Total (Rp)", "Lokasi", "Sumber Dana", "No. Dokumen",
                       "Kondisi", "Keterangan", "User"]
            widths = [5, 15, 12, 12, 28, 15, 14, 9, 9, 14, 14, 16, 14, 14, 12, 24, 18]
            fname = f"Barang-Masuk-{stamp}.xlsx"
        else:
            rows = [
                [
                    i + 1,
                    t.get("nomor", ""),
                    fdate(t.get("tanggal", "")),
                    t.get("kode_barang", ""),
                    t.get("nama_barang", ""),
                    t.get("kategori", ""),
                    "Barang Keluar" if jenis == "keluar" else "Penyesuaian",
                    t.get("jumlah", 0),
                    t.get("satuan", ""),
                    t.get("lokasi", ""),
                    t.get("tujuan", ""),
                    t.get("penerima", ""),
                    t.get("keperluan", ""),
                    t.get("keterangan", ""),
                    t.get("user_name", ""),
                ]
                for i, t in enumerate(docs)
            ]
            headers = ["No", "Nomor Transaksi", "Tanggal", "Kode Barang", "Nama Barang", "Kategori", "Jenis Transaksi",
                       "Jumlah", "Satuan", "Lokasi", "Tujuan", "Penerima", "Keperluan", "Keterangan", "User"]
            widths = [5, 15, 12, 12, 28, 15, 14, 9, 9, 16, 18, 20, 20, 24, 18]
            fname = f"Barang-{jenis.title()}-{stamp}.xlsx"
        _sheet(wb, f"Barang {jenis.title()}", kop1, kop2, headers, rows, widths)
        label = "Barang Masuk" if jenis == "masuk" else "Barang Keluar"

    elif jenis == "opname":
        docs = await db.stock_opnames.find().sort([("tanggal", -1), ("created_at", -1)]).to_list(5000)
        rows = []
        for o in docs:
            for d in o.get("details", []):
                selisih = int(d.get("selisih", 0))
                status = "Sesuai" if selisih == 0 else ("Lebih" if selisih > 0 else "Kurang")
                rows.append([
                    o.get("nomor", ""), fdate(o.get("tanggal", "")), o.get("lokasi", ""), o.get("kategori", ""),
                    d.get("kode", ""), d.get("nama", ""), d.get("satuan", ""), d.get("stok_sistem", 0),
                    d.get("stok_fisik", 0), selisih, status, d.get("keterangan", ""), o.get("petugas", ""),
                ])
        headers = ["No. Opname", "Tanggal", "Lokasi", "Kategori", "Kode", "Nama Barang", "Satuan",
                   "Stok Sistem", "Stok Fisik", "Selisih", "Status", "Keterangan", "Petugas"]
        _sheet(wb, "Stok Opname", kop1, kop2, headers, rows, [13, 12, 16, 15, 10, 28, 9, 11, 11, 9, 9, 26, 18])
        fname = f"Stok-Opname-{stamp}.xlsx"
        label = "Stok Opname"

    elif jenis == "kartu-stok":
        item_id = qp("item_id")
        if not item_id:
            raise HTTPException(status_code=400, detail="Pilih barang terlebih dahulu untuk export kartu stok")
        doc = await db.items.find_one({"id": item_id})
        if not doc:
            raise HTTPException(status_code=404, detail="Barang tidak ditemukan")
        txs = await db.item_transactions.find({"item_id": item_id, "status": "active"}).sort([("tanggal", 1), ("created_at", 1)]).to_list(100000)
        kartu = build_kartu(doc, txs, qp("awal"), qp("akhir"))
        rows = [["", "", "SALDO AWAL", "", "", "", kartu["saldo_awal"], ""]]
        rows += [
            [
                fdate(r["tanggal"]), r["nomor"], r["keterangan"], r["masuk"] or "", r["keluar"] or "",
                (f"+{r['penyesuaian']}" if r["penyesuaian"] > 0 else (str(r["penyesuaian"]) if r["penyesuaian"] else "")),
                r["saldo"], r["petugas"],
            ]
            for r in kartu["rows"]
        ]
        rows.append(["", "", "STOK AKHIR", kartu["total_masuk"], kartu["total_keluar"], kartu["total_penyesuaian"], kartu["saldo_akhir"], ""])
        _sheet(wb, "Kartu Stok", kop1, kop2,
               ["Tanggal", "No. Transaksi", "Keterangan", "Masuk", "Keluar", "Penyesuaian", "Saldo", "Petugas"],
               rows, [12, 15, 34, 9, 9, 12, 9, 18])
        fname = f"Kartu-Stok-{doc.get('kode', 'barang')}-{stamp}.xlsx"
        label = f"Kartu Stok {doc.get('nama', '')}"

    elif jenis == "template-import":
        rows = [
            ["", "Kertas A4 Sinar Dunia", "ATK", "Persediaan", "Sinar Dunia", "A4 70gr", "70 gram 500 lembar", "Rim", "2025", "BOS Reguler", "55000", "200", "Gudang Utama", "Baik", "", "", "Contoh baris - hapus sebelum mengisi"],
            ["", "Meja Siswa", "Meubelair", "Inventaris", "Lokal", "Kayu Jati", "Ukuran 60x40 cm", "Buah", "2024", "Komite", "350000", "120", "Ruang Kelas", "Baik", "", "", ""],
            ["", "Laptop Asus", "Alat Elektronik", "Inventaris/Aset", "Asus", "Vivobook 14", "Core i5 / RAM 8GB / SSD 512GB", "Unit", "2022", "Hibah", "7500000", "10", "Lab Komputer", "Baik", "SN-ASUS-001", "1", ""],
        ]
        _sheet(
            wb, "Data Barang", "TEMPLATE IMPORT DATA BARANG",
            "Isi data mulai baris ke-5. Lihat sheet 'Panduan' untuk aturan pengisian.",
            ["item_code", "item_name", "category", "item_type", "brand", "model", "specification", "unit",
             "acquisition_year", "funding_source", "unit_price", "initial_stock", "location", "condition",
             "serial_number", "nup", "description"],
            rows, [12, 30, 18, 14, 14, 16, 28, 9, 14, 14, 14, 12, 18, 12, 16, 8, 28],
        )
        panduan = [
            "PANDUAN IMPORT DATA BARANG",
            "1. Isi data pada sheet 'Data Barang' mulai baris ke-5 (jangan mengubah baris judul kolom).",
            "2. Kolom wajib: item_name (Nama Barang). Kolom lain boleh kosong (memakai nilai bawaan).",
            "3. item_type: 'Persediaan' atau 'Inventaris' (boleh juga 'Inventaris/Aset'). Kosong = Persediaan.",
            "4. category: ATK, Alat Kebersihan, Meubelair, Alat Elektronik, Peralatan Bengkel,",
            "   Peralatan Laboratorium, Peralatan Olahraga, atau Lainnya. Kosong = Lainnya.",
            "5. condition: Baik, Rusak Ringan, atau Rusak Berat. Kosong = Baik.",
            "6. acquisition_year, unit_price, dan initial_stock harus berupa angka (tanpa titik/koma ribuan).",
            "7. item_code: biarkan kosong untuk kode otomatis (BRG-xxxx).",
            "8. Jika item_code sudah ada di sistem, baris ditandai DUPLIKAT dan penanganannya mengikuti mode:",
            "   - Tambah Data Baru : baris ditolak (tidak membuat kode barang ganda)",
            "   - Update Data      : data barang diperbarui berdasarkan item_code",
            "   - Lewati Duplikat  : baris dilewati",
            "9. Simpan file sebagai .xlsx, .xls, atau .csv lalu unggah pada menu Import Excel.",
            "10. Periksa pratinjau hasil validasi, pilih mode, lalu tekan tombol Import.",
        ]
        _sheet(wb, "Panduan", "PANDUAN IMPORT", "", ["Petunjuk"], [[p] for p in panduan], [100])
        fname = "Template-Import-Data-Barang.xlsx"
        label = "Template Import"

    elif jenis == "pengguna":
        await require_admin(user)
        docs = await db.users.find().sort([("created_at", 1)]).to_list(1000)
        rows = [
            [u.get("nama", ""), u.get("username", ""), u.get("email", ""),
             "Administrator" if u.get("role") == "admin" else "Pengurus Barang",
             "Aktif" if u.get("aktif", True) else "Non-Aktif", u.get("nip", "")]
            for u in docs
        ]
        _sheet(wb, "Pengguna", kop1, kop2, ["Nama", "Username", "Email", "Role", "Status", "NIP"],
               rows, [26, 14, 26, 16, 11, 20])
        fname = f"Data-Pengguna-{stamp}.xlsx"
        label = "Data Pengguna"

    elif jenis == "audit":
        await require_admin(user)
        docs = await db.audit_logs.find().sort([("waktu", -1)]).to_list(1000)
        rows = [
            [
                d["waktu"].strftime("%d/%m/%Y %H:%M") if hasattr(d.get("waktu"), "strftime") else str(d.get("waktu", "")),
                d.get("user_name", ""), d.get("role", ""), d.get("aksi", ""), d.get("entitas", ""),
                d.get("entitas_id", ""), d.get("detail", ""), d.get("ip_address", ""),
            ]
            for d in docs
        ]
        _sheet(wb, "Riwayat Aktivitas", kop1, kop2,
               ["Waktu", "Pengguna", "Role", "Aksi", "Tabel", "Record ID", "Detail", "IP Address"],
               rows, [18, 22, 13, 14, 20, 30, 46, 14])
        fname = f"Riwayat-Aktivitas-{stamp}.xlsx"
        label = "Riwayat Aktivitas"

    else:
        raise HTTPException(status_code=404, detail="Jenis export tidak dikenal")

    await log_action(user, "EXPORT", "Sistem", detail=f"Export {label or jenis} (.xlsx)")
    return _response(wb, fname)
